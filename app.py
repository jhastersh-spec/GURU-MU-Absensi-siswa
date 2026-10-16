import base64
import hashlib
import io
import logging
import os
import secrets
import sqlite3
from datetime import date, datetime, timedelta, timezone
from functools import wraps
from pathlib import Path

import qrcode
from dotenv import load_dotenv
from flask import Flask, g, jsonify, render_template, request, send_file, session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfgen import canvas
from werkzeug.security import check_password_hash, generate_password_hash


ROOT = Path(__file__).resolve().parent
INSTANCE = ROOT / "instance"
load_dotenv(ROOT / ".env")
SCHEMA = """
CREATE TABLE IF NOT EXISTS users (
    id INTEGER PRIMARY KEY,
    username TEXT NOT NULL COLLATE NOCASE UNIQUE,
    password_hash TEXT NOT NULL,
    auth_version INTEGER NOT NULL DEFAULT 0,
    password_change_required INTEGER NOT NULL DEFAULT 0,
    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
);
CREATE TABLE IF NOT EXISTS students (
    id INTEGER PRIMARY KEY,
    name TEXT NOT NULL,
    nis TEXT NOT NULL UNIQUE,
    class_name TEXT NOT NULL,
    photo_url TEXT,
    qr_token TEXT NOT NULL UNIQUE,
    is_active INTEGER NOT NULL DEFAULT 1,
    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
    updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
);
CREATE TABLE IF NOT EXISTS attendance_sessions (
    id INTEGER PRIMARY KEY,
    attendance_date TEXT NOT NULL,
    class_name TEXT NOT NULL,
    is_open INTEGER NOT NULL DEFAULT 1,
    created_by INTEGER NOT NULL REFERENCES users(id),
    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
    closed_at TEXT,
    UNIQUE (attendance_date, class_name)
);
CREATE TABLE IF NOT EXISTS attendance (
    id INTEGER PRIMARY KEY,
    session_id INTEGER NOT NULL REFERENCES attendance_sessions(id),
    student_id INTEGER NOT NULL REFERENCES students(id),
    status TEXT NOT NULL CHECK (status IN ('HADIR', 'IZIN', 'SAKIT', 'ALPA', 'BELUM ABSEN')),
    scan_at TEXT,
    note TEXT,
    updated_by INTEGER NOT NULL REFERENCES users(id),
    updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
    UNIQUE (session_id, student_id)
);
CREATE TABLE IF NOT EXISTS attendance_audit_logs (
    id INTEGER PRIMARY KEY,
    session_id INTEGER NOT NULL REFERENCES attendance_sessions(id),
    student_id INTEGER,
    student_name TEXT NOT NULL,
    student_nis TEXT NOT NULL,
    previous_status TEXT,
    new_status TEXT NOT NULL,
    changed_by INTEGER NOT NULL REFERENCES users(id),
    changed_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
    note TEXT
);
CREATE TABLE IF NOT EXISTS login_attempts (
    ip_address TEXT PRIMARY KEY,
    attempts INTEGER NOT NULL DEFAULT 0,
    locked_until TEXT
);
CREATE TABLE IF NOT EXISTS login_rate_limits (
    scope_key TEXT PRIMARY KEY,
    attempts INTEGER NOT NULL DEFAULT 0,
    window_started_at TEXT NOT NULL,
    locked_until TEXT
);
CREATE INDEX IF NOT EXISTS idx_students_class ON students(class_name, is_active);
CREATE INDEX IF NOT EXISTS idx_attendance_session ON attendance(session_id, status);
CREATE INDEX IF NOT EXISTS idx_sessions_date ON attendance_sessions(attendance_date, class_name);
CREATE INDEX IF NOT EXISTS idx_audit_session ON attendance_audit_logs(session_id, changed_at);
"""
STATUSES = {"HADIR", "IZIN", "SAKIT", "ALPA", "BELUM ABSEN"}
FINAL_STATUSES = {"IZIN", "SAKIT", "ALPA"}


def password_is_strong(password):
    if len(password) < 12 or len(password) > 1024:
        return False
    character_groups = (
        any(character.islower() for character in password),
        any(character.isupper() for character in password),
        any(character.isdigit() for character in password),
        any(not character.isalnum() for character in password),
    )
    return sum(character_groups) >= 3


def _persistent_secret():
    INSTANCE.mkdir(parents=True, exist_ok=True)
    secret_file = INSTANCE / "secret_key"
    try:
        return secret_file.read_text(encoding="utf-8").strip()
    except FileNotFoundError:
        value = secrets.token_urlsafe(48)
        try:
            secret_file.write_text(value, encoding="utf-8")
            if os.name != "nt":
                secret_file.chmod(0o600)
        except FileExistsError:
            return secret_file.read_text(encoding="utf-8").strip()
        return value


def spreadsheet_safe(value):
    value = str(value or "")
    if value.lstrip(" \t\r\n").startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def create_app(test_config=None):
    application = Flask(__name__, instance_path=str(INSTANCE))
    configured_secret = os.environ.get("SECRET_KEY")
    if configured_secret and (
        len(configured_secret) < 32
        or configured_secret == "generate-a-random-secret-key"
    ):
        raise RuntimeError("SECRET_KEY harus berupa nilai acak rahasia minimal 32 karakter.")
    application.config.from_mapping(
        SECRET_KEY=configured_secret or _persistent_secret(),
        DATABASE_PATH=os.environ.get("DATABASE_PATH", str(INSTANCE / "attendance.db")),
        SESSION_COOKIE_SECURE=os.environ.get("COOKIE_SECURE", "1") == "1",
        SESSION_COOKIE_HTTPONLY=True,
        SESSION_COOKIE_SAMESITE="Lax",
        SESSION_COOKIE_NAME="guru_mu_session",
        PERMANENT_SESSION_LIFETIME=timedelta(hours=8),
        SESSION_REFRESH_EACH_REQUEST=False,
        MAX_CONTENT_LENGTH=2 * 1024 * 1024,
    )
    if test_config:
        application.config.update(test_config)
    Path(application.config["DATABASE_PATH"]).parent.mkdir(parents=True, exist_ok=True)

    def connect_db():
        connection = sqlite3.connect(application.config["DATABASE_PATH"], timeout=10)
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA foreign_keys = ON")
        connection.execute("PRAGMA journal_mode = WAL")
        return connection

    def initialize_database():
        connection = connect_db()
        try:
            connection.executescript(SCHEMA)
            columns = {row["name"] for row in connection.execute("PRAGMA table_info(users)")}
            if "auth_version" not in columns:
                connection.execute(
                    "ALTER TABLE users ADD COLUMN auth_version INTEGER NOT NULL DEFAULT 0"
                )
            if "password_change_required" not in columns:
                connection.execute(
                    "ALTER TABLE users ADD COLUMN password_change_required INTEGER NOT NULL DEFAULT 0"
                )
            connection.execute("UPDATE users SET password_change_required = 0")
            user_exists = connection.execute("SELECT 1 FROM users LIMIT 1").fetchone()
            if user_exists is None:
                username = os.environ.get("ADMIN_USERNAME", "").strip()
                password = os.environ.get("ADMIN_PASSWORD", "")
                if test_config:
                    username = test_config.get("ADMIN_USERNAME", username)
                    password = test_config.get("ADMIN_PASSWORD", password)
                if not username or not password or password == "change-this-to-a-long-unique-password":
                    raise RuntimeError(
                        "Akun admin belum tersedia. Atur ADMIN_USERNAME dan ADMIN_PASSWORD "
                        "di environment sebelum menjalankan aplikasi untuk pertama kali."
                    )
                if not password_is_strong(password):
                    raise RuntimeError(
                        "ADMIN_PASSWORD harus memiliki minimal 12 karakter dan mencakup "
                        "setidaknya tiga jenis karakter."
                    )
                connection.execute(
                    "INSERT INTO users (username, password_hash) VALUES (?, ?)",
                    (username, generate_password_hash(password)),
                )
            connection.commit()
        finally:
            connection.close()

    initialize_database()

    def get_db():
        if "db" not in g:
            g.db = connect_db()
        return g.db

    @application.teardown_appcontext
    def close_db(error=None):
        connection = g.pop("db", None)
        if connection is not None:
            if error is None:
                connection.commit()
            else:
                connection.rollback()
            connection.close()

    def json_error(message, status):
        return jsonify({"error": message}), status

    def json_object():
        data = request.get_json(silent=True)
        return data if isinstance(data, dict) else {}

    def login_required(function):
        @wraps(function)
        def wrapped(*args, **kwargs):
            user_id = session.get("user_id")
            if not user_id:
                return json_error("Silakan login kembali.", 401)
            user = get_db().execute(
                """SELECT auth_version, password_change_required
                   FROM users WHERE id = ?""",
                (user_id,),
            ).fetchone()
            if user is None or session.get("auth_version") != user["auth_version"]:
                session.clear()
                return json_error("Sesi login telah berakhir. Silakan login kembali.", 401)
            return function(*args, **kwargs)

        return wrapped

    def csrf_required(function):
        @wraps(function)
        def wrapped(*args, **kwargs):
            expected = session.get("csrf_token")
            supplied = request.headers.get("X-CSRF-Token", "")
            if not expected or not secrets.compare_digest(expected, supplied):
                return json_error("Permintaan tidak valid. Muat ulang halaman dan coba lagi.", 403)
            return function(*args, **kwargs)

        return wrapped

    def validate_date(value):
        try:
            return datetime.strptime(value, "%Y-%m-%d").date()
        except (TypeError, ValueError):
            raise ValueError("Tanggal harus menggunakan format YYYY-MM-DD.")

    def audit(connection, session_id, student, previous, new_status, note=None):
        connection.execute(
            """INSERT INTO attendance_audit_logs
               (session_id, student_id, student_name, student_nis, previous_status,
                new_status, changed_by, note)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                session_id,
                student["id"],
                student["name"],
                student["nis"],
                previous,
                new_status,
                session["user_id"],
                note,
            ),
        )

    def get_summary(day):
        db = get_db()
        counts = {status: 0 for status in STATUSES}
        rows = db.execute(
            """SELECT a.status, COUNT(DISTINCT a.student_id) AS total
               FROM attendance a
               JOIN attendance_sessions s ON s.id = a.session_id
               JOIN students st ON st.id = a.student_id
               WHERE s.attendance_date = ? AND st.is_active = 1
               GROUP BY a.status""",
            (day,),
        ).fetchall()
        for row in rows:
            counts[row["status"]] = row["total"]
        counts["total_students"] = db.execute(
            "SELECT COUNT(*) FROM students WHERE is_active = 1"
        ).fetchone()[0]
        counts["belum_absen"] = max(
            0, counts["total_students"] - sum(counts[s] for s in ("HADIR", "IZIN", "SAKIT", "ALPA"))
        )
        counts["persentase_kehadiran"] = round(
            counts["HADIR"] * 100 / counts["total_students"], 1
        ) if counts["total_students"] else 0
        return counts

    def report_rows(start, end, class_name="", search=""):
        clauses = ["s.attendance_date BETWEEN ? AND ?"]
        values = [start, end]
        if class_name:
            clauses.append("s.class_name = ?")
            values.append(class_name)
        if search:
            clauses.append("(st.name LIKE ? OR st.nis LIKE ?)")
            values.extend((f"%{search}%", f"%{search}%"))
        return get_db().execute(
            f"""SELECT st.id, st.name, st.nis, s.class_name, s.attendance_date,
                       a.status, a.scan_at
                FROM attendance a
                JOIN attendance_sessions s ON s.id = a.session_id
                JOIN students st ON st.id = a.student_id
                WHERE {' AND '.join(clauses)}
                ORDER BY s.attendance_date DESC, s.class_name, st.name""",
            values,
        ).fetchall()

    @application.after_request
    def security_headers(response):
        response.headers["X-Content-Type-Options"] = "nosniff"
        response.headers["X-Frame-Options"] = "DENY"
        response.headers["Referrer-Policy"] = "strict-origin-when-cross-origin"
        response.headers["Content-Security-Policy"] = (
            "default-src 'self'; img-src 'self' data: https:; style-src 'self'; "
            "script-src 'self'; connect-src 'self'; media-src 'self' blob:; "
            "object-src 'none'; frame-ancestors 'none'; base-uri 'self'; form-action 'self'"
        )
        response.headers["Permissions-Policy"] = "camera=(self), microphone=(), geolocation=()"
        response.headers["Cache-Control"] = "no-store"
        response.headers["Pragma"] = "no-cache"
        return response

    @application.get("/")
    def index():
        return render_template("index.html")

    @application.get("/api/auth")
    def auth_status():
        if "csrf_token" not in session:
            session["csrf_token"] = secrets.token_urlsafe(32)
        user = None
        if session.get("user_id"):
            user = get_db().execute(
                "SELECT id, username, auth_version FROM users WHERE id = ?",
                (session["user_id"],),
            ).fetchone()
            if user is None:
                session.clear()
                user = None
            elif session.get("auth_version") != user["auth_version"]:
                session.clear()
                user = None
        if "csrf_token" not in session:
            session["csrf_token"] = secrets.token_urlsafe(32)
        return jsonify(
            {
                "authenticated": user is not None,
                "username": user["username"] if user else None,
                "csrf_token": session["csrf_token"],
                "must_change_password": False,
            }
        )

    @application.post("/api/login")
    @csrf_required
    def login():
        data = json_object()
        username = str(data.get("username", "")).strip()
        password = str(data.get("password", ""))
        ip_address = request.remote_addr or "unknown"
        db = get_db()
        normalized_username = username.casefold()
        ip_digest = hashlib.sha256(ip_address.encode("utf-8")).hexdigest()
        identity_digest = hashlib.sha256(
            f"{ip_address}\0{normalized_username}".encode("utf-8")
        ).hexdigest()
        rate_keys = (f"ip:{ip_digest}", f"identity:{identity_digest}")
        now = datetime.now(timezone.utc)
        db.execute("BEGIN IMMEDIATE")
        for rate_key in rate_keys:
            attempt = db.execute(
                "SELECT locked_until FROM login_rate_limits WHERE scope_key = ?",
                (rate_key,),
            ).fetchone()
            if attempt and attempt["locked_until"]:
                locked_until = datetime.fromisoformat(attempt["locked_until"])
                if now < locked_until:
                    return json_error(
                        "Terlalu banyak percobaan login. Silakan coba lagi dalam 15 menit.",
                        429,
                    )
        user = db.execute(
            """SELECT id, username, password_hash, auth_version, password_change_required
               FROM users WHERE username = ? COLLATE NOCASE""",
            (username,),
        ).fetchone()
        if "dummy_password_hash" not in application.extensions:
            application.extensions["dummy_password_hash"] = generate_password_hash(
                secrets.token_urlsafe(32)
            )
        dummy_hash = application.extensions["dummy_password_hash"]
        password_too_long = len(password) > 1024
        hash_to_check = user["password_hash"] if user else dummy_hash
        password_matches = check_password_hash(
            hash_to_check,
            password if not password_too_long else password[:1024],
        )
        if (
            not username
            or len(username) > 120
            or password_too_long
            or user is None
            or not password_matches
        ):
            window_start = now - timedelta(minutes=15)
            db.execute(
                "DELETE FROM login_rate_limits WHERE window_started_at < ?",
                (window_start.isoformat(),),
            )
            attempt_count = 0
            for rate_key, limit in ((rate_keys[0], 15), (rate_keys[1], 5)):
                previous = db.execute(
                    """SELECT attempts, window_started_at FROM login_rate_limits
                       WHERE scope_key = ?""",
                    (rate_key,),
                ).fetchone()
                current_count = (
                    previous["attempts"] + 1
                    if previous and datetime.fromisoformat(previous["window_started_at"]) >= window_start
                    else 1
                )
                lock_until = now + timedelta(minutes=15) if current_count >= limit else None
                db.execute(
                    """INSERT INTO login_rate_limits
                       (scope_key, attempts, window_started_at, locked_until)
                       VALUES (?, ?, ?, ?)
                       ON CONFLICT(scope_key) DO UPDATE SET
                         attempts = excluded.attempts,
                         window_started_at = excluded.window_started_at,
                         locked_until = excluded.locked_until""",
                    (
                        rate_key,
                        current_count,
                        (previous["window_started_at"] if previous and current_count > 1 else now.isoformat()),
                        lock_until.isoformat() if lock_until else None,
                    ),
                )
                attempt_count = max(attempt_count, current_count)
            if attempt_count >= 5:
                return json_error(
                    "Terlalu banyak percobaan login. Silakan coba lagi dalam 15 menit.",
                    429,
                )
            return json_error("Username atau password salah.", 401)
        db.executemany(
            "DELETE FROM login_rate_limits WHERE scope_key = ?",
            ((rate_key,) for rate_key in rate_keys),
        )
        session.clear()
        session.permanent = True
        session["user_id"] = user["id"]
        session["auth_version"] = user["auth_version"]
        session["username"] = user["username"]
        session["csrf_token"] = secrets.token_urlsafe(32)
        return jsonify({
            "authenticated": True,
            "username": user["username"],
            "csrf_token": session["csrf_token"],
            "must_change_password": False,
        })

    @application.post("/api/logout")
    @csrf_required
    @login_required
    def logout():
        session.clear()
        session["csrf_token"] = secrets.token_urlsafe(32)
        return jsonify({"ok": True})

    @application.get("/api/dashboard")
    @login_required
    def dashboard():
        today = date.today().isoformat()
        db = get_db()
        sessions = db.execute(
            """SELECT id, attendance_date, class_name, is_open FROM attendance_sessions
               WHERE attendance_date = ? ORDER BY class_name""",
            (today,),
        ).fetchall()
        classes = [row["class_name"] for row in db.execute(
            "SELECT DISTINCT class_name FROM students WHERE is_active = 1 ORDER BY class_name"
        ).fetchall()]
        return jsonify({
            "date": today,
            "now": datetime.now().astimezone().isoformat(timespec="seconds"),
            "summary": get_summary(today),
            "classes": classes,
            "sessions": [dict(row) for row in sessions],
        })

    @application.get("/api/students")
    @login_required
    def list_students():
        search = request.args.get("search", "").strip()
        class_name = request.args.get("class", "").strip()
        clauses = ["is_active = 1"]
        values = []
        if search:
            clauses.append("(name LIKE ? OR nis LIKE ?)")
            values.extend((f"%{search}%", f"%{search}%"))
        if class_name:
            clauses.append("class_name = ?")
            values.append(class_name)
        students = get_db().execute(
            f"""SELECT id, name, nis, class_name, photo_url, created_at
                FROM students WHERE {' AND '.join(clauses)} ORDER BY class_name, name""",
            values,
        ).fetchall()
        return jsonify({"students": [dict(row) for row in students]})

    def validate_student(data):
        name = str(data.get("name", "")).strip()
        nis = str(data.get("nis", "")).strip()
        class_name = str(data.get("class_name", "")).strip()
        photo_url = str(data.get("photo_url", "")).strip() or None
        if not name or len(name) > 120:
            raise ValueError("Nama siswa wajib diisi (maksimal 120 karakter).")
        if not nis or len(nis) > 40:
            raise ValueError("NIS wajib diisi (maksimal 40 karakter).")
        if not class_name or len(class_name) > 60:
            raise ValueError("Kelas wajib diisi (maksimal 60 karakter).")
        if photo_url and not photo_url.startswith(("https://", "http://")):
            raise ValueError("Foto harus menggunakan alamat HTTP atau HTTPS.")
        return name, nis, class_name, photo_url

    @application.post("/api/students")
    @csrf_required
    @login_required
    def create_student():
        try:
            name, nis, class_name, photo_url = validate_student(json_object())
        except ValueError as error:
            return json_error(str(error), 400)
        db = get_db()
        try:
            cursor = db.execute(
                """INSERT INTO students (name, nis, class_name, photo_url, qr_token)
                   VALUES (?, ?, ?, ?, ?)""",
                (name, nis, class_name, photo_url, secrets.token_urlsafe(32)),
            )
            student = db.execute(
                """SELECT id, name, nis, class_name, photo_url, created_at
                   FROM students WHERE id = ?""",
                (cursor.lastrowid,),
            ).fetchone()
            open_sessions = db.execute(
                "SELECT id FROM attendance_sessions WHERE attendance_date = ? AND class_name = ? AND is_open = 1",
                (date.today().isoformat(), class_name),
            ).fetchall()
            for attendance_session in open_sessions:
                db.execute(
                    """INSERT INTO attendance (session_id, student_id, status, updated_by)
                       VALUES (?, ?, 'BELUM ABSEN', ?)""",
                    (attendance_session["id"], student["id"], session["user_id"]),
                )
                audit(db, attendance_session["id"], student, None, "BELUM ABSEN", "Siswa ditambahkan saat sesi dibuka.")
        except sqlite3.IntegrityError:
            return json_error("NIS tersebut sudah digunakan oleh siswa lain.", 409)
        return jsonify({"student": dict(student)}), 201

    @application.put("/api/students/<int:student_id>")
    @csrf_required
    @login_required
    def update_student(student_id):
        try:
            name, nis, class_name, photo_url = validate_student(json_object())
        except ValueError as error:
            return json_error(str(error), 400)
        db = get_db()
        student = db.execute(
            """SELECT id, name, nis, class_name, photo_url
               FROM students WHERE id = ? AND is_active = 1""",
            (student_id,),
        ).fetchone()
        if student is None:
            return json_error("Siswa tidak ditemukan.", 404)
        if class_name != student["class_name"]:
            existing_today = db.execute(
                """SELECT 1 FROM attendance a JOIN attendance_sessions s ON s.id = a.session_id
                   WHERE a.student_id = ? AND s.attendance_date = ? LIMIT 1""",
                (student_id, date.today().isoformat()),
            ).fetchone()
            if existing_today:
                return json_error("Kelas tidak dapat diubah setelah siswa memiliki catatan absensi hari ini.", 409)
            open_session = db.execute(
                """SELECT 1 FROM attendance_sessions
                   WHERE attendance_date = ? AND is_open = 1
                     AND class_name IN (?, ?) LIMIT 1""",
                (date.today().isoformat(), student["class_name"], class_name),
            ).fetchone()
            if open_session:
                return json_error("Kelas tidak dapat diubah saat sesi absensi kelas terkait masih terbuka.", 409)
        try:
            db.execute(
                """UPDATE students SET name = ?, nis = ?, class_name = ?, photo_url = ?,
                   updated_at = CURRENT_TIMESTAMP WHERE id = ?""",
                (name, nis, class_name, photo_url, student_id),
            )
        except sqlite3.IntegrityError:
            return json_error("NIS tersebut sudah digunakan oleh siswa lain.", 409)
        updated = db.execute(
            """SELECT id, name, nis, class_name, photo_url
               FROM students WHERE id = ?""",
            (student_id,),
        ).fetchone()
        return jsonify({"student": dict(updated)})

    @application.delete("/api/students/<int:student_id>")
    @csrf_required
    @login_required
    def delete_student(student_id):
        db = get_db()
        student = db.execute(
            "SELECT * FROM students WHERE id = ? AND is_active = 1", (student_id,)
        ).fetchone()
        if student is None:
            return json_error("Siswa tidak ditemukan.", 404)
        db.execute(
            "UPDATE students SET is_active = 0, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
            (student_id,),
        )
        return jsonify({"ok": True})

    @application.get("/api/students/<int:student_id>/card")
    @login_required
    def student_card(student_id):
        student = get_db().execute(
            """SELECT id, name, nis, class_name, photo_url, qr_token
               FROM students WHERE id = ? AND is_active = 1""",
            (student_id,),
        ).fetchone()
        if student is None:
            return json_error("Siswa tidak ditemukan.", 404)
        qr = qrcode.QRCode(box_size=7, border=2)
        qr.add_data(student["qr_token"])
        qr.make(fit=True)
        image = qr.make_image(fill_color="#102a43", back_color="white")
        output = io.BytesIO()
        image.save(output, format="PNG")
        data = base64.b64encode(output.getvalue()).decode("ascii")
        return jsonify({"student": {key: student[key] for key in ("id", "name", "nis", "class_name", "photo_url")},
                        "qr_data_url": f"data:image/png;base64,{data}"})

    @application.get("/api/sessions")
    @login_required
    def list_sessions():
        day = request.args.get("date", date.today().isoformat())
        try:
            validate_date(day)
        except ValueError as error:
            return json_error(str(error), 400)
        rows = get_db().execute(
            """SELECT s.id, s.attendance_date, s.class_name, s.is_open, s.created_at,
                      COUNT(a.id) AS total,
                      SUM(CASE WHEN a.status = 'HADIR' THEN 1 ELSE 0 END) AS hadir,
                      SUM(CASE WHEN a.status = 'IZIN' THEN 1 ELSE 0 END) AS izin,
                      SUM(CASE WHEN a.status = 'SAKIT' THEN 1 ELSE 0 END) AS sakit,
                      SUM(CASE WHEN a.status = 'ALPA' THEN 1 ELSE 0 END) AS alpa,
                      SUM(CASE WHEN a.status = 'BELUM ABSEN' THEN 1 ELSE 0 END) AS belum_absen
               FROM attendance_sessions s LEFT JOIN attendance a ON a.session_id = s.id
               WHERE s.attendance_date = ? GROUP BY s.id ORDER BY s.class_name""",
            (day,),
        ).fetchall()
        classes = [row["class_name"] for row in get_db().execute(
            "SELECT DISTINCT class_name FROM students WHERE is_active = 1 ORDER BY class_name"
        ).fetchall()]
        return jsonify({"sessions": [dict(row) for row in rows], "classes": classes})

    @application.post("/api/sessions")
    @csrf_required
    @login_required
    def open_session():
        data = json_object()
        class_name = str(data.get("class_name", "")).strip()
        day = str(data.get("date", date.today().isoformat()))
        try:
            validate_date(day)
        except ValueError as error:
            return json_error(str(error), 400)
        if not class_name or len(class_name) > 60:
            return json_error("Pilih kelas yang valid.", 400)
        db = get_db()
        existing = db.execute(
            "SELECT * FROM attendance_sessions WHERE attendance_date = ? AND class_name = ?",
            (day, class_name),
        ).fetchone()
        if existing and existing["is_open"]:
            return jsonify({"session": dict(existing)})
        with db:
            if existing:
                db.execute(
                    "UPDATE attendance_sessions SET is_open = 1, closed_at = NULL WHERE id = ?",
                    (existing["id"],),
                )
                session_id = existing["id"]
            else:
                cursor = db.execute(
                    """INSERT INTO attendance_sessions (attendance_date, class_name, created_by)
                       VALUES (?, ?, ?)""",
                    (day, class_name, session["user_id"]),
                )
                session_id = cursor.lastrowid
            students = db.execute(
                "SELECT * FROM students WHERE class_name = ? AND is_active = 1 ORDER BY name",
                (class_name,),
            ).fetchall()
            for student in students:
                inserted = db.execute(
                    """INSERT OR IGNORE INTO attendance (session_id, student_id, status, updated_by)
                       VALUES (?, ?, 'BELUM ABSEN', ?)""",
                    (session_id, student["id"], session["user_id"]),
                )
                if inserted.rowcount:
                    audit(db, session_id, student, None, "BELUM ABSEN", "Sesi absensi dibuka.")
        opened = db.execute("SELECT * FROM attendance_sessions WHERE id = ?", (session_id,)).fetchone()
        return jsonify({"session": dict(opened)}), 201

    @application.post("/api/sessions/<int:session_id>/close")
    @csrf_required
    @login_required
    def close_session(session_id):
        db = get_db()
        attendance_session = db.execute(
            "SELECT * FROM attendance_sessions WHERE id = ?", (session_id,)
        ).fetchone()
        if attendance_session is None:
            return json_error("Sesi tidak ditemukan.", 404)
        if attendance_session["is_open"]:
            db.execute(
                "UPDATE attendance_sessions SET is_open = 0, closed_at = CURRENT_TIMESTAMP WHERE id = ?",
                (session_id,),
            )
        return jsonify({"ok": True})

    @application.get("/api/sessions/<int:session_id>/roster")
    @login_required
    def session_roster(session_id):
        attendance_session = get_db().execute(
            "SELECT id, attendance_date, class_name, is_open FROM attendance_sessions WHERE id = ?",
            (session_id,),
        ).fetchone()
        if attendance_session is None:
            return json_error("Sesi tidak ditemukan.", 404)
        rows = get_db().execute(
            """SELECT st.id, st.name, st.nis, st.class_name, a.status, a.scan_at, a.note
               FROM students st
               LEFT JOIN attendance a ON a.student_id = st.id AND a.session_id = ?
               WHERE st.class_name = ? AND st.is_active = 1
               ORDER BY st.name""",
            (session_id, attendance_session["class_name"]),
        ).fetchall()
        return jsonify({"session": dict(attendance_session), "students": [dict(row) for row in rows]})

    @application.post("/api/scan")
    @csrf_required
    @login_required
    def scan_student():
        data = json_object()
        token = str(data.get("qr_token", "")).strip()
        try:
            session_id = int(data.get("session_id"))
        except (TypeError, ValueError):
            return json_error("Pilih sesi absensi terlebih dahulu.", 400)
        if not token or len(token) > 200:
            return json_error("Kode siswa tidak ditemukan.", 404)
        db = get_db()
        attendance_session = db.execute(
            "SELECT * FROM attendance_sessions WHERE id = ?", (session_id,)
        ).fetchone()
        if attendance_session is None:
            return json_error("Sesi absensi tidak ditemukan.", 404)
        if not attendance_session["is_open"]:
            return json_error("Sesi absensi sudah ditutup.", 409)
        student = db.execute(
            "SELECT * FROM students WHERE qr_token = ? AND is_active = 1", (token,)
        ).fetchone()
        if student is None:
            return json_error("Kode siswa tidak ditemukan.", 404)
        if student["class_name"] != attendance_session["class_name"]:
            return json_error("Siswa tidak termasuk dalam kelas sesi absensi ini.", 409)
        current = db.execute(
            "SELECT * FROM attendance WHERE session_id = ? AND student_id = ?",
            (session_id, student["id"]),
        ).fetchone()
        if current and current["status"] != "BELUM ABSEN":
            return json_error("Siswa sudah melakukan absensi hari ini.", 409)
        timestamp = datetime.now().astimezone().isoformat(timespec="seconds")
        if current:
            db.execute(
                """UPDATE attendance SET status = 'HADIR', scan_at = ?, note = NULL,
                   updated_by = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?""",
                (timestamp, session["user_id"], current["id"]),
            )
        else:
            db.execute(
                """INSERT INTO attendance (session_id, student_id, status, scan_at, updated_by)
                   VALUES (?, ?, 'HADIR', ?, ?)""",
                (session_id, student["id"], timestamp, session["user_id"]),
            )
        audit(db, session_id, student, current["status"] if current else None, "HADIR")
        return jsonify({
            "message": f"{student['name']} berhasil dicatat HADIR",
            "student": {"id": student["id"], "name": student["name"], "class_name": student["class_name"]},
            "scan_at": timestamp,
            "summary": get_summary(attendance_session["attendance_date"]),
        })

    @application.put("/api/sessions/<int:session_id>/students/<int:student_id>")
    @csrf_required
    @login_required
    def set_attendance(session_id, student_id):
        data = json_object()
        new_status = str(data.get("status", "")).upper()
        note = str(data.get("note", "")).strip() or None
        if new_status not in FINAL_STATUSES | {"BELUM ABSEN"}:
            return json_error("Status harus IZIN, SAKIT, ALPA, atau BELUM ABSEN.", 400)
        if note and len(note) > 500:
            return json_error("Catatan maksimal 500 karakter.", 400)
        db = get_db()
        attendance_session = db.execute(
            "SELECT * FROM attendance_sessions WHERE id = ?", (session_id,)
        ).fetchone()
        student = db.execute(
            "SELECT * FROM students WHERE id = ? AND is_active = 1", (student_id,)
        ).fetchone()
        if attendance_session is None or student is None:
            return json_error("Sesi atau siswa tidak ditemukan.", 404)
        if student["class_name"] != attendance_session["class_name"]:
            return json_error("Siswa tidak termasuk dalam kelas sesi ini.", 409)
        current = db.execute(
            "SELECT * FROM attendance WHERE session_id = ? AND student_id = ?",
            (session_id, student_id),
        ).fetchone()
        if current:
            db.execute(
                """UPDATE attendance SET status = ?, scan_at = NULL, note = ?,
                   updated_by = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?""",
                (new_status, note, session["user_id"], current["id"]),
            )
        else:
            db.execute(
                """INSERT INTO attendance (session_id, student_id, status, note, updated_by)
                   VALUES (?, ?, ?, ?, ?)""",
                (session_id, student_id, new_status, note, session["user_id"]),
            )
        audit(db, session_id, student, current["status"] if current else None, new_status, note)
        return jsonify({"ok": True, "summary": get_summary(attendance_session["attendance_date"])})

    @application.get("/api/audit")
    @login_required
    def audit_history():
        session_id = request.args.get("session_id", type=int)
        if session_id is None:
            return json_error("Pilih sesi absensi.", 400)
        rows = get_db().execute(
            """SELECT l.student_name, l.student_nis, l.previous_status, l.new_status,
                      l.changed_at, l.note, u.username
               FROM attendance_audit_logs l JOIN users u ON u.id = l.changed_by
               WHERE l.session_id = ? ORDER BY l.changed_at DESC, l.id DESC LIMIT 500""",
            (session_id,),
        ).fetchall()
        return jsonify({"logs": [dict(row) for row in rows]})

    @application.get("/api/reports")
    @login_required
    def attendance_report():
        period = request.args.get("period", "daily")
        if period not in {"daily", "weekly", "monthly"}:
            return json_error("Periode harus daily, weekly, atau monthly.", 400)
        try:
            if period == "monthly":
                month = request.args.get("month", date.today().strftime("%Y-%m"))
                month_date = datetime.strptime(month, "%Y-%m")
                start = month_date.date().replace(day=1)
                next_month = (start.replace(day=28) + timedelta(days=4)).replace(day=1)
                end = next_month - timedelta(days=1)
            elif period == "weekly":
                chosen = validate_date(request.args.get("date", date.today().isoformat()))
                start = chosen - timedelta(days=chosen.weekday())
                end = start + timedelta(days=6)
            else:
                chosen = validate_date(request.args.get("date", date.today().isoformat()))
                start = end = chosen
        except ValueError as error:
            return json_error(str(error), 400)
        class_name = request.args.get("class", "").strip()
        search = request.args.get("search", "").strip()
        records = report_rows(start.isoformat(), end.isoformat(), class_name, search)
        aggregate = {}
        for record in records:
            key = (record["id"], record["class_name"])
            item = aggregate.setdefault(key, {
                "id": record["id"], "name": record["name"], "nis": record["nis"],
                "class_name": record["class_name"], "HADIR": 0, "IZIN": 0, "SAKIT": 0,
                "ALPA": 0, "BELUM ABSEN": 0, "total": 0, "attendance_percentage": 0,
            })
            if record["status"] in STATUSES:
                item[record["status"]] += 1
                item["total"] += 1
        for item in aggregate.values():
            item["attendance_percentage"] = round(item["HADIR"] * 100 / item["total"], 1) if item["total"] else 0
        totals = {status: sum(item[status] for item in aggregate.values()) for status in STATUSES}
        return jsonify({
            "range": {"start": start.isoformat(), "end": end.isoformat()},
            "totals": totals,
            "students": list(aggregate.values()),
        })

    def export_parameters():
        period = request.args.get("period", "daily")
        if period not in {"daily", "weekly", "monthly"}:
            raise ValueError("Periode harus daily, weekly, atau monthly.")
        if period == "monthly":
            month_date = datetime.strptime(request.args.get("month", date.today().strftime("%Y-%m")), "%Y-%m")
            start = month_date.date().replace(day=1)
            end = (start.replace(day=28) + timedelta(days=4)).replace(day=1) - timedelta(days=1)
        elif period == "weekly":
            chosen = validate_date(request.args.get("date", date.today().isoformat()))
            start = chosen - timedelta(days=chosen.weekday())
            end = start + timedelta(days=6)
        else:
            start = end = validate_date(request.args.get("date", date.today().isoformat()))
        return start, end, report_rows(
            start.isoformat(), end.isoformat(),
            request.args.get("class", "").strip(), request.args.get("search", "").strip()
        )

    @application.get("/api/reports/export.xlsx")
    @login_required
    def export_excel():
        try:
            start, end, records = export_parameters()
        except ValueError as error:
            return json_error(str(error), 400)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Rekap Absensi"
        sheet.append(["Tanggal", "Nama", "NIS", "Kelas", "Status", "Waktu Scan"])
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="173B65")
        for item in records:
            values = [
                item["attendance_date"], item["name"], item["nis"], item["class_name"],
                item["status"], item["scan_at"] or "",
            ]
            sheet.append([spreadsheet_safe(value) for value in values])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column, width in zip("ABCDEF", (15, 32, 18, 16, 18, 30)):
            sheet.column_dimensions[column].width = width
        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)
        return send_file(output, as_attachment=True, download_name=f"rekap-absensi-{start}-{end}.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    @application.get("/api/reports/export.pdf")
    @login_required
    def export_pdf():
        try:
            start, end, records = export_parameters()
        except ValueError as error:
            return json_error(str(error), 400)
        output = io.BytesIO()
        page_size = landscape(A4)
        pdf = canvas.Canvas(output, pagesize=page_size)
        page_width, page_height = page_size

        def draw_header(page_number):
            pdf.setFillColor(colors.HexColor("#173B65"))
            pdf.rect(0, page_height - 58, page_width, 58, fill=1, stroke=0)
            pdf.setFillColor(colors.white)
            pdf.setFont("Helvetica-Bold", 17)
            pdf.drawString(36, page_height - 37, "GURU MU - Rekap Absensi")
            pdf.setFont("Helvetica", 10)
            pdf.drawRightString(page_width - 36, page_height - 37, f"{start} sampai {end}")
            pdf.setFillColor(colors.HexColor("#102A43"))
            pdf.setFont("Helvetica-Bold", 9)
            headers = [("Tanggal", 36), ("Nama", 130), ("NIS", 350), ("Kelas", 455), ("Status", 555), ("Waktu Scan", 650)]
            y = page_height - 83
            for title, x in headers:
                pdf.drawString(x, y, title)
            pdf.setStrokeColor(colors.HexColor("#D8E2EC"))
            pdf.line(36, y - 7, page_width - 36, y - 7)
            return y - 25, page_number

        y, page_number = draw_header(1)
        for record in records:
            if y < 42:
                pdf.showPage()
                page_number += 1
                y, _ = draw_header(page_number)
            values = [
                record["attendance_date"], record["name"][:34], record["nis"],
                record["class_name"], record["status"], (record["scan_at"] or "")[:24],
            ]
            for text, x in zip(values, (36, 130, 350, 455, 555, 650)):
                pdf.setFillColor(colors.HexColor("#243B53"))
                pdf.setFont("Helvetica", 8)
                pdf.drawString(x, y, str(text))
            y -= 17
        pdf.save()
        output.seek(0)
        return send_file(output, as_attachment=True, download_name=f"rekap-absensi-{start}-{end}.pdf",
                         mimetype="application/pdf")

    @application.post("/api/settings/password")
    @csrf_required
    @login_required
    def change_password():
        data = json_object()
        current_password = str(data.get("current_password", ""))
        new_password = str(data.get("new_password", ""))
        if not password_is_strong(new_password):
            return json_error(
                "Password baru harus memiliki minimal 12 karakter dan mencakup "
                "setidaknya tiga jenis karakter (huruf kecil, huruf besar, angka, simbol).",
                400,
            )
        db = get_db()
        user = db.execute(
            "SELECT password_hash FROM users WHERE id = ?", (session["user_id"],)
        ).fetchone()
        if user is None or not check_password_hash(user["password_hash"], current_password):
            return json_error("Password saat ini salah.", 400)
        if check_password_hash(user["password_hash"], new_password):
            return json_error("Password baru harus berbeda dari password saat ini.", 400)
        new_auth_version = session["auth_version"] + 1
        db.execute(
            """UPDATE users SET password_hash = ?, auth_version = ?,
               password_change_required = 0 WHERE id = ?""",
            (generate_password_hash(new_password), new_auth_version, session["user_id"]),
        )
        session["auth_version"] = new_auth_version
        session["csrf_token"] = secrets.token_urlsafe(32)
        return jsonify({
            "ok": True,
            "csrf_token": session["csrf_token"],
            "must_change_password": False,
        })

    return application


app = create_app()

if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO)
    app.run(host=os.environ.get("HOST", "127.0.0.1"), port=int(os.environ.get("PORT", "5000")), debug=False)
